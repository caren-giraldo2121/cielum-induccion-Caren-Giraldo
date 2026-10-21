"""
Bot RPA - Reporte diario de pedidos (proyecto Cielum, parte 7)

Entra al frontend React como un usuario de prueba, filtra los pedidos del dia,
lee la tabla pagina por pagina y genera reporte_diario.xlsx.

Uso:
    python reporte_diario.py                 # pedidos de hoy
    python reporte_diario.py --fecha 2026-10-21
    python reporte_diario.py --headless      # sin abrir ventana

NOTA: esto es RPA a proposito (ejercicio). En la vida real se usaria la API
GET /pedidos?desde=...&hasta=...  -> ver README.md de esta carpeta.
"""
import argparse
import os
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

CARPETA = Path(__file__).resolve().parent
CAPTURAS = CARPETA / "capturas"
SALIDA = CARPETA / "reporte_diario.xlsx"

load_dotenv(CARPETA / ".env")


def leer_config():
    p = argparse.ArgumentParser(description="Genera reporte_diario.xlsx leyendo el frontend")
    p.add_argument("--fecha", default=os.getenv("RPA_FECHA") or date.today().isoformat(),
                   help="Fecha a reportar (YYYY-MM-DD). Por defecto hoy.")
    p.add_argument("--headless", action="store_true",
                   default=os.getenv("RPA_HEADLESS", "false").lower() == "true")
    args = p.parse_args()
    datetime.strptime(args.fecha, "%Y-%m-%d")  # valida el formato
    return {
        "url": os.getenv("RPA_URL", "http://localhost:5173").rstrip("/"),
        "usuario": os.getenv("RPA_USUARIO", "caren@cielum.com"),
        "password": os.getenv("RPA_PASSWORD", "123456"),
        "fecha": args.fecha,
        "headless": args.headless,
    }


def es_respuesta_pedidos(resp):
    return "/api/pedidos" in resp.url and resp.request.method == "GET"


def iniciar_sesion(page, cfg):
    print(f"[1/4] Abriendo {cfg['url']}/login")
    page.goto(f"{cfg['url']}/login")
    page.fill("#email", cfg["usuario"])
    page.fill("#password", cfg["password"])
    page.click("button[type=submit]")
    try:
        page.wait_for_url("**/pedidos", timeout=10_000)
    except PlaywrightTimeout:
        alerta = page.locator(".alert.error")
        motivo = alerta.inner_text() if alerta.count() else "no redirigio a /pedidos"
        raise RuntimeError(f"No se pudo iniciar sesion: {motivo}")
    print("      Sesion iniciada")


def esperar_resultado(page):
    """Espera a que termine de cargar: o hay tabla, o aparece el estado vacio, o un error."""
    page.locator("table").or_(page.locator(".estado-vacio")).or_(page.locator(".card .alert.error")).first.wait_for(timeout=15_000)
    error = page.locator(".card .alert.error")
    if error.count():
        raise RuntimeError(f"El frontend mostro un error: {error.first.inner_text()}")


def filtrar_dia(page, fecha):
    print(f"[2/4] Filtrando pedidos del {fecha}")
    # Esperar a que la primera carga termine antes de tocar filtros
    esperar_resultado(page)
    page.fill("input[name=desde]", fecha)
    page.fill("input[name=hasta]", fecha)
    page.select_option("select[name=limite]", "50")
    with page.expect_response(es_respuesta_pedidos, timeout=15_000):
        page.get_by_role("button", name="Buscar").click()
    esperar_resultado(page)


def leer_tabla(page):
    print("[3/4] Leyendo la tabla de pedidos")
    pedidos = []
    pagina = 1
    while True:
        if page.locator(".estado-vacio").count():
            break
        filas = page.locator("table tbody tr")
        for i in range(filas.count()):
            celdas = filas.nth(i).locator("td")
            pedidos.append({
                "id_pedido": int(celdas.nth(0).inner_text().strip()),
                "id_cliente": int(celdas.nth(1).inner_text().strip()),
                "fecha": celdas.nth(2).inner_text().strip(),
                "estado": celdas.nth(3).inner_text().strip(),
            })
        print(f"      Pagina {pagina}: {filas.count()} pedidos")

        siguiente = page.get_by_role("button", name="Siguiente →")
        if siguiente.count() == 0 or siguiente.is_disabled():
            break
        with page.expect_response(es_respuesta_pedidos, timeout=15_000):
            siguiente.click()
        esperar_resultado(page)
        pagina += 1
    return pedidos


def generar_excel(pedidos, cfg):
    print(f"[4/4] Generando {SALIDA.name}")
    wb = Workbook()
    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E78")

    ws = wb.active
    ws.title = "Pedidos"
    columnas = ["ID pedido", "ID cliente", "Fecha", "Estado"]
    ws.append(columnas)
    for p in pedidos:
        ws.append([p["id_pedido"], p["id_cliente"], p["fecha"], p["estado"]])
    for col, ancho in zip("ABCD", [12, 12, 14, 16]):
        ws.column_dimensions[col].width = ancho
    for celda in ws[1]:
        celda.font, celda.fill, celda.alignment = encabezado, relleno, Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    if pedidos:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(pedidos) + 1}"

    rs = wb.create_sheet("Resumen")
    rs.append(["Reporte diario de pedidos - Cielum"])
    rs["A1"].font = Font(bold=True, size=14)
    rs.append(["Fecha reportada", cfg["fecha"]])
    rs.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    rs.append(["Fuente", f"{cfg['url']}/pedidos (bot RPA, usuario {cfg['usuario']})"])
    rs.append(["Total pedidos", len(pedidos)])
    rs.append([])
    rs.append(["Estado", "Cantidad"])
    for celda in rs[rs.max_row]:
        celda.font, celda.fill = encabezado, relleno
    for estado, cantidad in sorted(Counter(p["estado"] for p in pedidos).items()):
        rs.append([estado, cantidad])
    rs.column_dimensions["A"].width = 20
    rs.column_dimensions["B"].width = 55

    wb.save(SALIDA)


def main():
    cfg = leer_config()
    CAPTURAS.mkdir(exist_ok=True)
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(headless=cfg["headless"], slow_mo=0 if cfg["headless"] else 150)
        page = navegador.new_page(viewport={"width": 1366, "height": 850})
        try:
            iniciar_sesion(page, cfg)
            filtrar_dia(page, cfg["fecha"])
            page.screenshot(path=str(CAPTURAS / "rpa-pedidos-del-dia.png"), full_page=True)
            pedidos = leer_tabla(page)
            generar_excel(pedidos, cfg)
            print(f"\nListo: {len(pedidos)} pedidos -> {SALIDA}")
        except Exception as e:
            page.screenshot(path=str(CAPTURAS / "rpa-error.png"), full_page=True)
            print(f"\nERROR: {e}\nCaptura del momento del error: {CAPTURAS / 'rpa-error.png'}", file=sys.stderr)
            sys.exit(1)
        finally:
            navegador.close()


if __name__ == "__main__":
    main()
